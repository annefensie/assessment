#!/usr/bin/env python3
"""
extract_framework.py — Extract assessment_framework_expansion_updated.xlsx
into per-sheet CSV and JSON files in ../data/, and regenerate
../src/FrameworkData.gs with embedded JS constants.

Usage:
    python3 scripts/extract_framework.py path/to/assessment_framework_expansion_updated.xlsx
"""

import sys
import json
import csv
import os
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

REPO_ROOT = Path(__file__).parent.parent
DATA_DIR = REPO_ROOT / "data"
SRC_DIR = REPO_ROOT / "src"

SHEET_VAR_MAP = {
    "Verbs":                   "VERBS",
    "Behaviors":               "BEHAVIORS",
    "Tasks":                   "TASKS",
    "Deliverables":            "DELIVERABLES",
    "Discipline_Map":          "DISCIPLINE_MAP",
    "AI_Resistance":           "AI_RESISTANCE",
    "Sources":                 "SOURCES",
    "Program_Disciplines":     "PROGRAM_DISCIPLINES",
    "App_Implementation_Notes":"APP_IMPL_NOTES",
}


def sheet_to_dicts(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [
        str(h).strip() if h is not None else f"col_{i}"
        for i, h in enumerate(rows[0])
    ]
    result = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        d = {h: (str(v).strip() if v is not None else "") for h, v in zip(headers, row)}
        result.append(d)
    return result


def write_csv(path, rows):
    if not rows:
        return
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)


def write_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def generate_framework_data_gs(all_data, out_path):
    from datetime import date
    lines = [
        "// AUTO-GENERATED from assessment_framework_expansion_updated.xlsx",
        "// Regenerate with: python3 scripts/extract_framework.py <path/to/xlsx>",
        f"// Last generated: {date.today().isoformat()}",
        "",
        "var FRAMEWORK_DATA = (function() {",
        "  'use strict';",
        "",
    ]
    for sheet, varname in SHEET_VAR_MAP.items():
        if sheet not in all_data:
            continue
        jstr = json.dumps(all_data[sheet], ensure_ascii=False)
        lines.append(f"  var {varname} = {jstr};")
        lines.append("")
    lines += [
        "  return {",
    ]
    for sheet, varname in SHEET_VAR_MAP.items():
        if sheet in all_data:
            lines.append(f"    {varname}: {varname},")
    lines += [
        "  };",
        "})();",
    ]
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


def main():
    if len(sys.argv) < 2:
        print(f"Usage: {sys.argv[0]} path/to/assessment_framework_expansion_updated.xlsx")
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    if not xlsx_path.exists():
        print(f"File not found: {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    DATA_DIR.mkdir(exist_ok=True)

    print(f"Loading {xlsx_path}...")
    wb = openpyxl.load_workbook(xlsx_path)
    all_data = {}

    for name in wb.sheetnames:
        rows = sheet_to_dicts(wb[name])
        all_data[name] = rows
        short = name.lower().replace(" ", "_")
        write_csv(DATA_DIR / f"{short}.csv", rows)
        write_json(DATA_DIR / f"{short}.json", rows)
        print(f"  {name}: {len(rows)} rows")

    write_json(DATA_DIR / "framework_all.json", all_data)
    print(f"\nWrote {DATA_DIR / 'framework_all.json'}")

    gs_path = SRC_DIR / "FrameworkData.gs"
    generate_framework_data_gs(all_data, gs_path)
    print(f"Wrote {gs_path}")

    print("\nDone.")


if __name__ == "__main__":
    main()
